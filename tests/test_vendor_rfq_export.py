"""Capture log #12 — one RFQ workbook per vendor per show, in Larry's own
template columns — and note 2 steps 4–5 (the vendor on every surface), plus
the dormant Google Sheets path (#13).

Jason, 2026-09-06: RFQ, per vendor per show, match the template, the GHC26
Local Labor RFQ PDF is the shape; departments as the catalogue names them;
xlsx now, the Sheets button built but hidden until Google is configured.
"""
import datetime as dt
import io
import json

import openpyxl
import pytest

import google_sheets
import rfq_export


D0 = dt.date(2026, 10, 24)


def _fixture(db):
    """GHC26 in miniature: two days, riggers from VRA, lighting from Sparks,
    a scenic line with no vendor set."""
    from models import (Show, ScheduleDay, ScheduleActivity, CrewRow, Company,
                        Position, ShowDepartmentVendor)
    show = Show(name="Grace Hopper RFQ", code="GHCRFQ")
    db.session.add(show); db.session.flush()
    vra = Company(name="VRA RFQ", code="VRA", contact_name="Vince", email="v@vra.test",
                  ot_after_hours=8, dt_after_hours=12)
    sparks = Company(name="Sparks RFQ", code="SPK")
    db.session.add_all([vra, sparks]); db.session.flush()
    rig = Position(title="Rigger High RFQ", department="Rigging", is_local_labor=True)
    hand = Position(title="Lighting Hand RFQ", department="Lighting", is_local_labor=True)
    scen = Position(title="Scenic Hand RFQ", department="Scenic", is_local_labor=True)
    db.session.add_all([rig, hand, scen]); db.session.flush()
    db.session.add_all([
        ShowDepartmentVendor(show_id=show.id, department="Rigging", company_id=vra.id, job_number="VRA-77"),
        ShowDepartmentVendor(show_id=show.id, department="Lighting", company_id=sparks.id),
    ])
    for i, (time, lines) in enumerate([
        ("08:00", [(rig, 4, 11), (hand, 9, 11), (scen, 2, 8)]),
        ("07:00", [(rig, 3, 12), (hand, 5, 12)]),
    ]):
        day = ScheduleDay(show_id=show.id, date=D0 + dt.timedelta(days=i),
                          label=f"Load-in day {i + 1}", phase="Load In")
        db.session.add(day); db.session.flush()
        act = ScheduleActivity(day_id=day.id, time=time, description="CREW START", sort_order=1)
        db.session.add(act); db.session.flush()
        for n, (pos, qty, hrs) in enumerate(lines, start=2):
            db.session.add(CrewRow(activity_id=act.id, position_id=pos.id, position=pos.title,
                                   qty=qty, hours=hrs, crew_type="Local Crew", sort_order=n,
                                   task="Overnight pre-rig" if pos is rig and i == 0 else None))
    db.session.commit()
    return show, vra, sparks


# ── the split by vendor ──────────────────────────────────────────────────────

def test_lines_go_to_their_departments_vendor_and_the_rest_to_unassigned(app, db):
    show, vra, sparks = _fixture(db)
    with app.app_context():
        buckets = rfq_export.lines_by_vendor(show)
    assert list(buckets) == [sparks.id, vra.id, rfq_export.UNASSIGNED]   # name order, Unassigned last
    assert [l["qty"] for l in buckets[vra.id]["lines"]] == [4, 3]
    assert buckets[vra.id]["job_number"] == "VRA-77"
    assert [l["position"] for l in buckets[rfq_export.UNASSIGNED]["lines"]] == ["Scenic Hand RFQ"]
    s = rfq_export.summary(buckets[sparks.id])
    assert (s["lines"], s["assignments"], s["person_hours"], s["days"]) == (2, 14, 9 * 11 + 5 * 12, 2)


# ── the workbook ─────────────────────────────────────────────────────────────

def _wb(app, db, show, bucket):
    with app.app_context():
        wb = rfq_export.build_workbook(show, bucket, issue_date=dt.date(2026, 9, 6))
    return openpyxl.load_workbook(io.BytesIO(rfq_export.workbook_bytes(wb)))


def _day(v):
    """openpyxl hands dates back as datetimes."""
    return v.date() if isinstance(v, dt.datetime) else v


def test_workbook_has_the_templates_tabs_and_columns(app, db):
    show, vra, sparks = _fixture(db)
    with app.app_context():
        bucket = rfq_export.lines_by_vendor(show)[vra.id]
    wb = _wb(app, db, show, bucket)
    assert wb.sheetnames == ["00 RFQ Setup", "02 Daily Summary", "03 Labor Schedule", "07 Lists"]
    ws = wb["03 Labor Schedule"]
    assert [c.value for c in ws[6]] == rfq_export.LABOR_HEADERS
    # two rigging lines, day 1 and day 2
    assert ws["A7"].value == 1 and ws["A8"].value == 2
    assert _day(ws["B7"].value) == D0
    assert ws["E7"].value == "8:00 AM" and ws["F7"].value == "7:00 PM"     # 08:00 + 11h
    assert ws["G7"].value == "Rigging"
    assert ws["H7"].value == "Rigger High RFQ"
    assert ws["I7"].value == "Load-in day 1 — Overnight pre-rig"
    assert ws["J7"].value == 4
    # VRA's own terms: OT after 8, DT after 12 -> 11h = 8 + 3 + 0
    assert (ws["K7"].value, ws["L7"].value, ws["M7"].value) == (8.0, 3.0, None)
    # vendor prices; the sheet computes
    assert ws["N7"].value is None
    assert ws["O7"].value == '=IF(N7="","",N7*1.5)'
    assert "J7*(" in ws["Q7"].value
    # the template's colours: ADI yellow on the ask, blue on the formulas
    assert ws["J7"].fill.fgColor.rgb.endswith("FFF2CC")
    assert ws["Q7"].fill.fgColor.rgb.endswith("EAF3F8")


def test_setup_sheet_names_the_show_the_vendor_and_the_counts(app, db):
    show, vra, sparks = _fixture(db)
    with app.app_context():
        bucket = rfq_export.lines_by_vendor(show)[vra.id]
    ws = _wb(app, db, show, bucket)["00 RFQ Setup"]
    cells = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=3).value
             for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value}
    assert cells["RFQ ID"] == "RFQ-GHCRFQ-VRA"
    assert cells["Opportunity / Show ID"] == "GHCRFQ"
    assert cells["RFQ Discipline"] == "Local Labor"
    assert _day(cells["Issue Date"]) == dt.date(2026, 9, 6)
    assert cells["Event / Program"] == "Grace Hopper RFQ"
    assert cells["Vendor Company"] == "VRA RFQ"
    assert cells["Vendor Contact"] == "Vince"
    assert cells["Vendor Job Number"] == "VRA-77"
    assert _day(cells["Master Labor Window Start (Travel In)"]) == D0
    assert cells["Dated call lines"] == 2
    assert cells["Assignments (bodies)"] == 7
    assert cells["Planning person-hours"] == 4 * 11 + 3 * 12


def test_daily_summary_matches_the_ghc26_cover(app, db):
    show, vra, sparks = _fixture(db)
    with app.app_context():
        bucket = rfq_export.lines_by_vendor(show)[sparks.id]
    ws = _wb(app, db, show, bucket)["02 Daily Summary"]
    assert [c.value for c in ws[4]] == ["Date", "Phase", "Assignments", "Person-Hours", "Primary Activity"]
    assert (_day(ws["A5"].value), ws["C5"].value, ws["D5"].value, ws["E5"].value) == (D0, 9, 99.0, "Load-in day 1")
    assert (ws["C6"].value, ws["D6"].value) == (5, 60.0)


def test_lists_sheet_carries_the_catalogue(app, db):
    show, vra, sparks = _fixture(db)
    with app.app_context():
        bucket = rfq_export.lines_by_vendor(show)[vra.id]
    ws = _wb(app, db, show, bucket)["07 Lists"]
    depts = [ws.cell(row=r, column=1).value for r in range(5, 20)]
    assert depts[:3] == ["General", "Rigging", "Lighting"]
    positions = [ws.cell(row=r, column=2).value for r in range(5, 60)]
    assert "Rigger High RFQ" in positions


def test_unassigned_workbook_says_so(app, db):
    show, vra, sparks = _fixture(db)
    with app.app_context():
        bucket = rfq_export.lines_by_vendor(show)[rfq_export.UNASSIGNED]
    ws = _wb(app, db, show, bucket)["00 RFQ Setup"]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "no vendor assigned" in text
    assert rfq_export.filename_for(show, bucket, dt.date(2026, 9, 6)) == \
        "GHCRFQ_Local_Labor_RFQ_Unassigned_20260906.xlsx"


# ── the routes ───────────────────────────────────────────────────────────────

def test_picker_lists_a_card_per_vendor_and_downloads_a_workbook(client, db):
    show, vra, sparks = _fixture(db)
    html = client.get(f"/shows/{show.id}/rfq").get_data(as_text=True)
    assert "VRA RFQ" in html and "Sparks RFQ" in html and "Unassigned vendor" in html
    assert "Job # VRA-77" in html
    assert "To Google Sheets" not in html          # not configured
    assert "Connect Google" not in html
    r = client.get(f"/shows/{show.id}/rfq/{vra.id}.xlsx")
    assert r.status_code == 200
    assert "GHCRFQ_Local_Labor_RFQ_VRA_" in r.headers["Content-Disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    assert wb["03 Labor Schedule"]["H7"].value == "Rigger High RFQ"
    assert client.get(f"/shows/{show.id}/rfq/Unassigned.xlsx").status_code == 200
    assert client.get(f"/shows/{show.id}/rfq/999999.xlsx").status_code == 404


def test_the_door_is_on_the_schedule_overview_and_the_show_page(client, db):
    show, vra, sparks = _fixture(db)
    assert f"/shows/{show.id}/rfq" in client.get(f"/shows/{show.id}/schedule").get_data(as_text=True)
    assert f"/shows/{show.id}/rfq" in client.get(f"/shows/{show.id}").get_data(as_text=True)


# ── note 2 steps 4-5: the vendor on every surface ───────────────────────────

def test_the_vendor_shows_on_the_day_page_department_header(client, db):
    show, vra, sparks = _fixture(db)
    day = show.days[0]
    html = client.get(f"/shows/{show.id}/schedule/{day.id}").get_data(as_text=True)
    assert "· VRA RFQ" in html
    assert "· Sparks RFQ" in html


def test_the_vendor_rides_the_local_labor_line_into_the_oss_exports(app, client, db):
    show, vra, sparks = _fixture(db)
    from oss_export import build_master_items
    with app.app_context():
        items, _ = build_master_items(show, [], [])
    labels = [l["label"] for it in items for l in (it.get("local_lines") or [])]
    assert any(l.endswith("— VRA RFQ") for l in labels)
    assert any(l.endswith("— Sparks RFQ") for l in labels)
    assert any("Scenic Hand RFQ" in l and "—" not in l.split("×")[-1] for l in labels)
    # and the show book header
    book = client.get(f"/shows/{show.id}/oss/show-book").get_data(as_text=True)
    assert "· VRA RFQ" in book


# ── Google: dormant until configured ────────────────────────────────────────

def test_google_routes_are_404_until_configured(client, db, monkeypatch):
    monkeypatch.delenv("GOOGLE_OAUTH_CLIENT_ID", raising=False)
    monkeypatch.delenv("GOOGLE_OAUTH_CLIENT_SECRET", raising=False)
    assert google_sheets.is_configured() is False
    assert client.get("/google/connect").status_code == 404
    assert client.get("/google/callback?code=x").status_code == 404
    show, vra, sparks = _fixture(db)
    assert client.post(f"/shows/{show.id}/rfq/{vra.id}/sheet").status_code == 404


def test_configured_picker_offers_connect_then_the_button(client, db, monkeypatch):
    monkeypatch.setenv("GOOGLE_OAUTH_CLIENT_ID", "id.apps.test")
    monkeypatch.setenv("GOOGLE_OAUTH_CLIENT_SECRET", "s3cr3t")
    show, vra, sparks = _fixture(db)
    html = client.get(f"/shows/{show.id}/rfq").get_data(as_text=True)
    assert "Connect Google" in html
    assert 'disabled title="Connect Google first"' in html
    r = client.get("/google/connect")
    assert r.status_code == 302
    assert r.location.startswith(google_sheets.AUTH_URL)
    assert "access_type=offline" in r.location
    from models import GoogleCredential
    GoogleCredential.replace("refresh-token-1", "larry@adiexpgroup.com")
    html = client.get(f"/shows/{show.id}/rfq").get_data(as_text=True)
    assert "connected as <strong>larry@adiexpgroup.com" in html
    assert "disabled" not in html.split("To Google Sheets")[0][-300:]


def test_publish_rfq_walks_the_drive_api(monkeypatch):
    """token -> 01 - RFQs -> <show> folder -> upload as Sheet -> share."""
    calls = []

    def fake_http(method, url, data=None, headers=None, content_type=None):
        calls.append((method, url))
        if url == google_sheets.TOKEN_URL:
            return {"access_token": "at"}
        if method == "GET" and url.startswith(google_sheets.DRIVE_FILES + "?"):
            # first lookup finds the root folder, second finds nothing
            return {"files": [{"id": "root1", "name": "01 - RFQs"}]} if "01+-+RFQs" in url else {"files": []}
        if method == "POST" and url == google_sheets.DRIVE_FILES + "?fields=id":
            assert data["parents"] == ["root1"]
            return {"id": "sub1"}
        if url.startswith(google_sheets.DRIVE_UPLOAD):
            assert headers["Authorization"] == "Bearer at"
            assert content_type.startswith("multipart/related")
            assert b"application/vnd.google-apps.spreadsheet" in data
            return {"id": "file1", "name": "GHC_RFQ", "webViewLink": "https://docs.google.com/x"}
        if url.endswith("/file1/permissions"):
            assert data == {"type": "anyone", "role": "writer"}
            return {"id": "perm"}
        raise AssertionError(url)

    monkeypatch.setattr(google_sheets, "_http", fake_http)
    monkeypatch.setenv("GOOGLE_OAUTH_CLIENT_ID", "id")
    monkeypatch.setenv("GOOGLE_OAUTH_CLIENT_SECRET", "s")
    created = google_sheets.publish_rfq("rt", "GHC26", "GHC_RFQ.xlsx", b"PK\x03\x04")
    assert created["webViewLink"] == "https://docs.google.com/x"
    methods = [m for m, _ in calls]
    assert methods == ["POST", "GET", "GET", "POST", "POST", "POST"]
